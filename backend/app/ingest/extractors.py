"""TextExtractor (B7, PROF-2, PROF-5): file bytes -> plain text, failing loudly.

Type is decided by magic bytes, never by the filename or the client's content-type.
Caps: size, pages, DOCX uncompressed size and entry count (zip bombs), parse timeout.
LaTeX is read as text and never compiled (PDF-3). No OCR (PROF-5).
Adding a format = one strategy + one dict entry (OCP).
"""

import asyncio
import re
import zipfile
from dataclasses import dataclass
from io import BytesIO
from typing import Protocol

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.errors import InvalidInput, UnreadableUpload

MIN_TEXT_CHARS = 200  # less than this from a PDF = image-only scan
_SPACES = re.compile(r"[ \t]+")
_BLANK_LINES = re.compile(r"\n{3,}")


def _tidy(text: str) -> str:
    """Layout-mode text pads columns with spaces: one space between words, at most one
    blank line between blocks."""
    lines = (_SPACES.sub(" ", line).strip() for line in text.splitlines())
    return _BLANK_LINES.sub("\n\n", "\n".join(lines)).strip()


@dataclass(frozen=True, slots=True)
class Extracted:
    source_type: str   # pdf | docx | latex
    ext: str
    text: str


class Extractor(Protocol):
    source_type: str
    ext: str

    def extract(self, data: bytes) -> str: ...


class PdfExtractor:
    source_type, ext = "pdf", "pdf"

    def __init__(self, max_pages: int) -> None:
        self._max_pages = max_pages

    def extract(self, data: bytes) -> str:
        try:
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise UnreadableUpload("This PDF is password-protected.", reason="encrypted")
            if len(reader.pages) > self._max_pages:
                raise UnreadableUpload(f"CVs up to {self._max_pages} pages, please.",
                                       reason="too_many_pages")
            # Layout mode places glyphs by position. The default mode drops the space at a
            # font switch ("through **Python**" -> "throughPython") and splits kerned pairs
            # ("T echnology"); the parser copies verbatim, so those would reach the CV.
            text = "\n".join(_tidy(page.extract_text(extraction_mode="layout") or "")
                             for page in reader.pages)
        except PdfReadError as exc:
            raise UnreadableUpload("This PDF looks corrupt.", reason="corrupt_pdf") from exc
        if len(text.strip()) < MIN_TEXT_CHARS:
            raise UnreadableUpload("This PDF has no selectable text (a scan?). Upload a text "
                                   "PDF or DOCX.", reason="image_only")
        return text


class DocxExtractor:
    source_type, ext = "docx", "docx"

    def __init__(self, max_uncompressed: int) -> None:
        self._max_uncompressed = max_uncompressed

    def extract(self, data: bytes) -> str:
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                infos = archive.infolist()
                if len(infos) > 1000 or sum(i.file_size for i in infos) > self._max_uncompressed:
                    raise UnreadableUpload("This file is too large inside.", reason="zip_bomb")
            document = Document(BytesIO(data))
        except (zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise UnreadableUpload("This DOCX looks corrupt.", reason="corrupt_docx") from exc
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(line for line in lines if line.strip())
        if len(text.strip()) < MIN_TEXT_CHARS:
            raise UnreadableUpload("This DOCX has almost no text.", reason="empty_docx")
        return text


class LatexTextExtractor:
    """Pasted LaTeX source: content only. Comments dropped; never compiled."""

    source_type, ext = "latex", "tex"
    _COMMENT = re.compile(r"(?<!\\)%.*$", re.MULTILINE)

    def extract(self, data: bytes) -> str:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise UnreadableUpload("LaTeX must be UTF-8 text.", reason="encoding") from exc
        text = self._COMMENT.sub("", text)
        if len(text.strip()) < MIN_TEXT_CHARS:
            raise UnreadableUpload("This LaTeX has almost no content.", reason="empty_latex")
        return text[:40000]


class TextExtractor:
    def __init__(self, max_bytes: int, max_pages: int, max_uncompressed: int,
                 timeout_s: int) -> None:
        self._max_bytes = max_bytes
        self._timeout_s = timeout_s
        self._by_type: dict[str, Extractor] = {
            "pdf": PdfExtractor(max_pages),
            "docx": DocxExtractor(max_uncompressed),
            "latex": LatexTextExtractor(),
        }

    @staticmethod
    def detect(data: bytes) -> str:
        if data.startswith(b"%PDF-"):
            return "pdf"
        if data.startswith(b"PK\x03\x04") and b"word/" in data[:4096]:
            return "docx"
        raise UnreadableUpload("Upload a PDF or DOCX file.", reason="unsupported_type")

    async def extract(self, data: bytes, source_type: str | None = None) -> Extracted:
        if not data:
            raise InvalidInput("The file is empty.", reason="empty_file")
        if len(data) > self._max_bytes:
            raise InvalidInput(f"Files up to {self._max_bytes // (1024 * 1024)} MB, please.",
                               reason="too_large")
        kind = source_type or self.detect(data)
        extractor = self._by_type[kind]
        try:
            # Parsing is CPU-bound and libraries can stall on hostile files: run off the event
            # loop with a deadline.
            async with asyncio.timeout(self._timeout_s):
                text = await asyncio.to_thread(extractor.extract, data)
        except TimeoutError as exc:
            raise UnreadableUpload("This file took too long to read.", reason="timeout") from exc
        return Extracted(extractor.source_type, extractor.ext, text)
